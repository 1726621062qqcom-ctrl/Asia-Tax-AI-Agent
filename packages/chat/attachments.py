"""Read user documents as untrusted text in a bounded, killable worker."""
from __future__ import annotations
import asyncio
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile
from xml.etree import ElementTree as ET
from packages.chat.pdf_attachments import MAX_BYTES, MAX_TEXT, extract_pdf

ROOT = Path(__file__).resolve().parents[2]
EXTENSIONS = {'.pdf', '.doc', '.docx', '.xls', '.xlsx', '.png', '.jpg', '.jpeg', '.webp'}
MARKER = '\n\n[Taxrora attachment]\n'


def checked_zip(data):
    archive = zipfile.ZipFile(BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 2000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
        archive.close()
        raise ValueError('attachment_expanded_limit')
    return archive


def result(kind, chunks, warnings=None):
    if not chunks or not any(c['text'].strip() for c in chunks):
        raise ValueError('attachment_no_text')
    if sum(len(c['text']) + len(c.get('label', '')) for c in chunks) > MAX_TEXT:
        raise ValueError('pdf_too_much_text')
    return {'kind': kind, 'page_count': len(chunks), 'pages': chunks, 'empty_pages': [], 'warnings': warnings or []}


def word_xml(data):
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with checked_zip(data) as z:
        root = ET.fromstring(z.read('word/document.xml'))
    body = root.find(ns + 'body')
    if body is None:
        raise ValueError('attachment_invalid')
    def paragraph(p):
        return ''.join(n.text or '' if n.tag == ns+'t' else '\t' if n.tag == ns+'tab' else '\n'
                       for n in p.iter() if n.tag in {ns+'t', ns+'tab', ns+'br'})
    lines = []
    for element in body:
        if element.tag == ns+'p':
            lines.append(paragraph(element))
        elif element.tag == ns+'tbl':
            lines.append('[Table]')
            for row in element.findall(ns+'tr'):
                lines.append('\t'.join(' / '.join(paragraph(p) for p in cell.findall('.//'+ns+'p')) for cell in row.findall(ns+'tc')))
    return result('word', [{'page': 1, 'label': 'Document body / tables', 'text': '\n'.join(lines).strip()}], ['word_body_only'])


def word_doc(data):
    if not data.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        raise ValueError('attachment_invalid')
    env = dict(os.environ)
    if sys.platform == 'darwin':
        command = ['/usr/bin/textutil', '-convert', 'txt', '-format', 'doc', '-stdin', '-stdout']
    else:
        bundled = ROOT / '.tools/usr/bin/antiword'
        binary = shutil.which('antiword') or (str(bundled) if bundled.exists() else None)
        if not binary:
            raise ValueError('attachment_doc_unavailable')
        if binary == str(bundled):
            env['ANTIWORDHOME'] = str(ROOT / '.tools/usr/share/antiword')
        command = [binary, '-m', 'UTF-8.txt', '-']
    completed = subprocess.run(command, input=data, capture_output=True, timeout=12, env=env)
    if completed.returncode:
        raise ValueError('attachment_invalid')
    text = completed.stdout.decode('utf-8', errors='replace').strip()
    return result('word', [{'page': 1, 'label': 'Document text', 'text': text}], ['word_body_only'])


def excel_xlsx(data):
    from openpyxl import load_workbook
    with checked_zip(data):
        pass
    book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        if len(book.worksheets) > 20:
            raise ValueError('attachment_sheet_limit')
        chunks, total, cells, has_formulas = [], 0, 0, False
        for number, sheet in enumerate(book.worksheets, 1):
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 100000:
                raise ValueError('attachment_cell_limit')
            lines = []
            for row, values in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                parts = []
                for cell, value in zip(row, values):
                    if cell.value is None:
                        continue
                    cells += 1
                    if cells > 10000:
                        raise ValueError('attachment_cell_limit')
                    content = str(cell.value)
                    if cell.data_type == 'f':
                        has_formulas = True
                        content += ' [cached value: ' + (str(value.value) if value.value is not None else 'unavailable') + ']'
                    if cell.number_format != 'General':
                        content += f' [format: {cell.number_format}]'
                    parts.append(f'{cell.coordinate}: {content}')
                line = ' | '.join(parts)
                total += len(line)
                if total > MAX_TEXT:
                    raise ValueError('pdf_too_much_text')
                if line:
                    lines.append(line)
            if lines:
                chunks.append({'page': number, 'label': f'Worksheet: {sheet.title}', 'text': '\n'.join(lines)})
        return result('excel', chunks, ['excel_values_only'] + (['excel_cached_formulas'] if has_formulas else []))
    finally:
        book.close(); cached.close()


def excel_xls(data):
    import xlrd
    from openpyxl.utils import get_column_letter
    book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    try:
        if book.nsheets > 20:
            raise ValueError('attachment_sheet_limit')
        chunks, cells, total = [], 0, 0
        for number, sheet in enumerate(book.sheets(), 1):
            if sheet.nrows * sheet.ncols > 100000:
                raise ValueError('attachment_cell_limit')
            lines = []
            for r in range(sheet.nrows):
                parts = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    cells += 1
                    if cells > 10000:
                        raise ValueError('attachment_cell_limit')
                    value = xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat() if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                    fmt = book.format_map[book.xf_list[sheet.cell_xf_index(r, c)].format_key].format_str
                    parts.append(f'{get_column_letter(c+1)}{r+1}: {value}' + (f' [format: {fmt}]' if fmt != 'General' else ''))
                line = ' | '.join(parts)
                total += len(line)
                if total > MAX_TEXT:
                    raise ValueError('pdf_too_much_text')
                if line: lines.append(line)
            if lines: chunks.append({'page': number, 'label': f'Worksheet: {sheet.name}', 'text': '\n'.join(lines)})
        return result('excel', chunks, ['excel_values_only', 'excel_cached_formulas'])
    finally:
        book.release_resources()


def image_document(data):
    import base64
    import warnings
    from PIL import Image, ImageOps
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(BytesIO(data)) as original:
            if original.format not in ('PNG', 'JPEG', 'WEBP'):
                raise ValueError('attachment_invalid')
            if original.width * original.height > 20_000_000:
                raise ValueError('attachment_image_dimensions')
            original.load()
            picture = ImageOps.exif_transpose(original).convert('RGB')
    thumb = picture.copy()
    thumb.thumbnail((200, 160))
    buffer = BytesIO()
    thumb.save(buffer, format='JPEG', quality=55)
    if len(buffer.getvalue()) > 9000:
        thumb.thumbnail((100, 80)); buffer = BytesIO(); thumb.save(buffer, format='JPEG', quality=40)
    thumbnail = 'data:image/jpeg;base64,' + base64.b64encode(buffer.getvalue()).decode('ascii')
    picture.thumbnail((2400, 2400))
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
    reader = RapidOCR(intra_op_num_threads=1, inter_op_num_threads=1)
    found, _ = reader(np.asarray(picture))
    text = '\n'.join(item[1] for item in (found or []) if item[2] >= .5)
    payload = result('image', [{'page': 1, 'label': 'Image OCR (unverified)', 'text': text}], ['image_ocr_only'])
    payload['thumbnail'] = thumbnail
    return payload


def model_attachment_text(text):
    """Display thumbnails are not model input; only OCR text is provided."""
    if MARKER not in text:
        return text
    prefix, raw = text.split(MARKER, 1)
    try:
        payload = json.loads(raw)
        payload.pop('thumbnail', None)
        return prefix + MARKER + json.dumps(payload, ensure_ascii=False)
    except (ValueError, TypeError, AttributeError):
        return text


def extract_document(data, extension):
    if len(data) > MAX_BYTES: raise ValueError('pdf_too_large')
    if extension not in EXTENSIONS: raise ValueError('attachment_type')
    try:
        if extension in ('.png', '.jpg', '.jpeg', '.webp'): return image_document(data)
        if extension == '.pdf': return {**extract_pdf(data), 'kind': 'pdf', 'warnings': []}
        return {'.docx': word_xml, '.doc': word_doc, '.xlsx': excel_xlsx, '.xls': excel_xls}[extension](data)
    except Exception as exc:
        code = str(exc)
        if code.startswith(('pdf_', 'attachment_')): raise ValueError(code) from None
        raise ValueError('attachment_invalid') from None


async def read_document(data, extension):
    proc = await asyncio.create_subprocess_exec(sys.executable, '-m', 'packages.chat.attachments', extension,
        cwd=str(ROOT), start_new_session=True, stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.DEVNULL)
    try:
        output, _ = await asyncio.wait_for(proc.communicate(data), 45 if extension in ('.png', '.jpg', '.jpeg', '.webp') else 20)
        if proc.returncode: raise ValueError('attachment_invalid')
        payload = json.loads(output)
        if 'error' in payload: raise ValueError(payload['error'])
        return payload
    except asyncio.TimeoutError:
        raise ValueError('pdf_timeout') from None
    finally:
        if proc.returncode is None:
            import signal
            try: os.killpg(proc.pid, signal.SIGKILL)
            except ProcessLookupError: pass
            await proc.wait()


if __name__ == '__main__':
    import resource
    os.environ['OMP_NUM_THREADS'] = '1'
    os.environ['OPENBLAS_NUM_THREADS'] = '1'
    image_input = sys.argv[1] in ('.png', '.jpg', '.jpeg', '.webp')
    resource.setrlimit(resource.RLIMIT_CPU, (35, 35) if image_input else (15, 15))
    if sys.platform == 'linux': resource.setrlimit(resource.RLIMIT_AS, ((1536 if image_input else 384) * 1024 * 1024,) * 2)
    try: payload = extract_document(sys.stdin.buffer.read(MAX_BYTES + 1), sys.argv[1])
    except Exception as exc: payload = {'error': str(exc) if str(exc).startswith(('pdf_', 'attachment_')) else 'attachment_invalid'}
    print(json.dumps(payload, ensure_ascii=False))


def attachment_source(text):
    """Return readable page text, never metadata or thumbnails, for fact checks."""
    from packages.chat.pdf_attachments import MARKER as LEGACY_MARKER
    for marker in (MARKER, LEGACY_MARKER):
        if marker not in text:
            continue
        prefix, raw = text.split(marker, 1)
        try:
            payload = json.loads(raw)
            pages = payload['pages']
            if not isinstance(pages, list) or not pages or not all(isinstance(p, dict) and isinstance(p.get('text'), str) for p in pages):
                return None
            return prefix + '\n\n[User-provided attachment text; unverified]\n' + '\n\n'.join(p['text'] for p in pages)
        except (ValueError, TypeError, KeyError):
            return None
    return None
