import asyncio
from io import BytesIO
from pathlib import Path
import shutil
import sys
import zipfile
import pytest
from openpyxl import Workbook
from packages.chat.attachments import extract_document, read_document

FIXTURES = Path(__file__).parent / 'fixtures'


def docx():
    out = BytesIO()
    with zipfile.ZipFile(out, 'w') as z:
        z.writestr('word/document.xml', '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Hong Kong company</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>Amount</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>2000000</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>''')
    return out.getvalue()


def xlsx():
    b = Workbook(); s = b.active; s.title = 'Dividends'; s['A1'] = 2000000; s['B1'] = '=A1*0.05'; s['C1'] = .05; s['C1'].number_format = '0%'
    b.create_sheet('Company')['A1'] = '香港公司'
    out = BytesIO(); b.save(out); return out.getvalue()


def test_docx_keeps_body_and_table_order():
    r = asyncio.run(read_document(docx(), '.docx'))
    assert r['kind'] == 'word'
    assert 'Hong Kong company\n[Table]\nAmount\t2000000' in r['pages'][0]['text']
    assert r['warnings'] == ['word_body_only']


def test_xlsx_keeps_sheets_cells_formats_and_does_not_evaluate_formulas():
    r = asyncio.run(read_document(xlsx(), '.xlsx'))
    assert r['kind'] == 'excel' and len(r['pages']) == 2
    assert r['pages'][0]['label'] == 'Worksheet: Dividends'
    assert 'A1: 2000000' in r['pages'][0]['text']
    assert 'B1: =A1*0.05 [cached value: unavailable]' in r['pages'][0]['text']
    assert 'C1: 0.05 [format: 0%]' in r['pages'][0]['text']
    assert '香港公司' in r['pages'][1]['text']
    assert 'excel_cached_formulas' in r['warnings']


def test_xls_reads_real_binary_workbook():
    r = asyncio.run(read_document((FIXTURES/'attachment.xls').read_bytes(), '.xls'))
    assert r['kind'] == 'excel'
    assert 'A2: 2000000' in r['pages'][0]['text']


@pytest.mark.skipif(sys.platform != 'darwin' and not shutil.which('antiword'), reason='legacy converter not installed')
def test_doc_reads_real_binary_word():
    r = asyncio.run(read_document((FIXTURES/'attachment.doc').read_bytes(), '.doc'))
    assert 'CNY 2000000' in r['pages'][0]['text']


@pytest.mark.parametrize('ext', ['.doc', '.docx', '.xls', '.xlsx'])
def test_invalid_file_is_not_read_as_text(ext):
    with pytest.raises(ValueError, match='attachment_invalid'):
        extract_document(b'not an office document', ext)


def test_sheet_and_expanded_zip_limits():
    b = Workbook()
    for i in range(20): b.create_sheet(str(i))
    out=BytesIO();b.save(out)
    with pytest.raises(ValueError, match='attachment_sheet_limit'): extract_document(out.getvalue(), '.xlsx')
    out=BytesIO()
    with zipfile.ZipFile(out, 'w') as z:
        for i in range(2001): z.writestr(str(i), '')
    with pytest.raises(ValueError, match='attachment_expanded_limit'): extract_document(out.getvalue(), '.docx')

@pytest.mark.parametrize('name,data', [('case.docx', docx()), ('case.xlsx', xlsx()), ('case.xls', (FIXTURES/'attachment.xls').read_bytes())])
def test_attachment_api_accepts_office_documents(name, data):
    from fastapi.testclient import TestClient
    from apps.api.main import app
    from apps.api.auth import optional_user
    app.dependency_overrides[optional_user] = lambda: {'id': 'office-test'}
    try:
        with TestClient(app) as client:
            response = client.post('/api/attachments', params={'filename':name}, content=data, headers={'x-asiatax-request':'1'})
            assert response.status_code == 200, response.text
            assert response.json()['name'] == name
            assert response.json()['pages']
    finally:
        app.dependency_overrides.clear()
